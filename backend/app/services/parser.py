import os
import hashlib
from typing import List, Dict, Any
import fitz  # PyMuPDF
from docx import Document as DocxDocument

class ParsedDocument:
    def __init__(self, filename: str, file_hash: str, page_count: int, pages: List[Dict[str, Any]], full_text: str):
        self.filename = filename
        self.file_hash = file_hash
        self.page_count = page_count
        self.pages = pages  # List of {"page_number": int, "text": str, "sections": List[str]}
        self.full_text = full_text

class DocumentParser:
    @staticmethod
    def compute_sha256(file_bytes: bytes) -> str:
        return hashlib.sha256(file_bytes).hexdigest()

    @classmethod
    def parse(cls, file_bytes: bytes, filename: str) -> ParsedDocument:
        if not file_bytes:
            raise ValueError(f"Empty document: {filename} contains 0 bytes.")

        file_hash = cls.compute_sha256(file_bytes)
        ext = os.path.splitext(filename)[1].lower()

        if ext == ".pdf":
            return cls._parse_pdf(file_bytes, filename, file_hash)
        elif ext == ".docx":
            return cls._parse_docx(file_bytes, filename, file_hash)
        elif ext in [".txt", ".md", ".markdown"]:
            return cls._parse_text(file_bytes, filename, file_hash)
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Allowed: .pdf, .docx, .txt, .md")

    @classmethod
    def _parse_pdf(cls, file_bytes: bytes, filename: str, file_hash: str) -> ParsedDocument:
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception as e:
            raise ValueError(f"Corrupted or invalid PDF file '{filename}': {str(e)}")

        pages: List[Dict[str, Any]] = []
        full_text_parts: List[str] = []

        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_text = page.get_text("text") or ""
            sections = cls._extract_sections_from_text(page_text)
            pages.append({
                "page_number": page_idx + 1,
                "text": page_text,
                "sections": sections
            })
            full_text_parts.append(page_text)

        full_text = "\n\n".join(full_text_parts)
        if not full_text.strip():
            raise ValueError(f"PDF document '{filename}' contains no extractable text.")

        return ParsedDocument(
            filename=filename,
            file_hash=file_hash,
            page_count=max(len(pages), 1),
            pages=pages,
            full_text=full_text
        )

    @classmethod
    def _parse_docx(cls, file_bytes: bytes, filename: str, file_hash: str) -> ParsedDocument:
        import io
        try:
            doc = DocxDocument(io.BytesIO(file_bytes))
        except Exception as e:
            raise ValueError(f"Corrupted or invalid DOCX file '{filename}': {str(e)}")

        full_text_parts: List[str] = []
        sections: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                full_text_parts.append(text)
                if p.style and "heading" in p.style.name.lower():
                    sections.append(text)

        full_text = "\n\n".join(full_text_parts)
        if not full_text.strip():
            raise ValueError(f"DOCX document '{filename}' contains no extractable text.")

        # Estimate pages based on word count (~400 words per page)
        words = len(full_text.split())
        page_count = max(1, (words + 399) // 400)

        # Distribute into synthetic page chunks
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        chunk_size = max(1, len(paragraphs) // page_count) if paragraphs else 1
        pages: List[Dict[str, Any]] = []

        for p_idx in range(page_count):
            start = p_idx * chunk_size
            end = start + chunk_size if p_idx < page_count - 1 else len(paragraphs)
            p_text = "\n\n".join(paragraphs[start:end])
            pages.append({
                "page_number": p_idx + 1,
                "text": p_text,
                "sections": cls._extract_sections_from_text(p_text) or sections
            })

        return ParsedDocument(
            filename=filename,
            file_hash=file_hash,
            page_count=page_count,
            pages=pages,
            full_text=full_text
        )

    @classmethod
    def _parse_text(cls, file_bytes: bytes, filename: str, file_hash: str) -> ParsedDocument:
        try:
            full_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                full_text = file_bytes.decode("latin-1")
            except Exception as e:
                raise ValueError(f"Unable to decode text document '{filename}': {str(e)}")

        if not full_text.strip():
            raise ValueError(f"Text document '{filename}' is empty.")

        sections = cls._extract_sections_from_text(full_text)
        words = len(full_text.split())
        page_count = max(1, (words + 399) // 400)

        lines = [l for l in full_text.splitlines() if l.strip()]
        chunk_size = max(1, len(lines) // page_count)
        pages: List[Dict[str, Any]] = []

        for p_idx in range(page_count):
            start = p_idx * chunk_size
            end = start + chunk_size if p_idx < page_count - 1 else len(lines)
            p_text = "\n".join(lines[start:end])
            pages.append({
                "page_number": p_idx + 1,
                "text": p_text,
                "sections": cls._extract_sections_from_text(p_text)
            })

        return ParsedDocument(
            filename=filename,
            file_hash=file_hash,
            page_count=page_count,
            pages=pages,
            full_text=full_text
        )

    @staticmethod
    def _extract_sections_from_text(text: str) -> List[str]:
        sections: List[str] = []
        for line in text.splitlines():
            s = line.strip()
            if s.startswith("#") or (s.isupper() and len(s) > 3 and len(s) < 80):
                clean = s.lstrip("#").strip()
                if clean and clean not in sections:
                    sections.append(clean)
        return sections
